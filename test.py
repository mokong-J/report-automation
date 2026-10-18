from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def write_df(ws, df):
    # 写入表头
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", start_color="4472C4")
        cell.alignment = Alignment(horizontal="center")

    # 写入数据
    # 获取每行数据
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        # 获取每列数据
        for col_idx, col_value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=col_value)


file_path = r"E:\PythonProject\项目练习\01_excel_marge\cleaned_sales_data.xlsx"

df = pd.read_excel(file_path)

wb = Workbook()
ws1 = wb.active
ws1.title = "原始数据"

write_df(ws1, df)

# 制作透视表
df["月份"] = df["Date"].str[:7]
pivot = pd.pivot_table(
    df,
    index="月份",
    columns="Region",
    values="Amount",
    aggfunc="sum",
    fill_value=0,  # 空值补 0
    margins=True,  # 自动加合计行/列
    margins_name="Total",
)
pivot = pivot.reset_index()

# 写入透视表
ws2 = wb.create_sheet("透视表")
write_df(ws2, pivot)

# 图表
ws3 = wb.create_sheet("图表")

date_str = datetime.now().strftime("%Y%m%d")

output_path = Path(__file__).parent / "output" / f"销售周报_{date_str}.xlsx"
try:
    wb.save(output_path)
except PermissionError:
    print("Excel 已打开，请关闭后重试")
